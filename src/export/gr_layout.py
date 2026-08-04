"""Pixel-faithful Maharashtra GR renderer — the downloaded draft must look
exactly like an original government GR.

Layout is calibrated span-by-span against the gold-standard specimen
``data/raw_pdfs/202402271742452721.pdf`` (School Education & Sports Dept,
3 pages): US-Letter page, 72/69 pt margins, two line grids (19.3 pt tight /
22.2 pt body), top-right bold subject block, centred bold letterhead, hanging
वाचा and प्रति lists, justified body, running header with a thin-over-thick
rule and a right-aligned "पृष्ठ N पैकी n" footer on pages 2+.

Two invariants the rest of the codebase depends on:

* The canonical draft text stays ASCII digits end-to-end — the review
  pipeline, scorecard regexes and digit audits never see Devanagari digits.
  Conversion to ०१२३४५६७८९ happens at render time only, inside this module.
* A draft must NOT fabricate a digital-signature stamp. The stamp is applied
  at issuance; a genuine pre-issuance draft ends at the ( name ) +
  designation block, and so does this renderer.
"""
import re
import sys
from datetime import date as date_cls
from pathlib import Path

from fpdf import FPDF

ROOT = Path(__file__).resolve().parents[2]
FONT_DIR = ROOT / "assets" / "fonts"
sys.path.append(str(ROOT / "src" / "draftgen"))
from generate import MARATHI_MONTHS, financial_year, unique_code  # noqa: E402

# ---- geometry, all in pt, measured from data/raw_pdfs/202402271742452721.pdf ----
PAGE_W, PAGE_H = 612.0, 792.0      # US Letter — real GRs are NOT A4
MARGIN_L = 72.0
MARGIN_R = 69.0                    # text right edge = 543
BODY_W = PAGE_W - MARGIN_L - MARGIN_R          # 471
CONTENT_TOP = 49.5                 # first text line top on every page
BREAK_MARGIN = 36.0                # auto-page-break trigger

LINE_H_BODY = 22.2                 # प्रस्तावना / clauses pitch
LINE_H_TIGHT = 19.3                # subject, letterhead, वाचा, प्रति pitch
PARA_GAP = 10.0                    # extra space between body paragraphs
SECTION_GAP = 10.0                 # before a bold section label
FIRST_INDENT = 36.0                # preamble first-line indent

SIZE_BODY = 12.0
SIZE_TITLE = 13.0                  # "महाराष्ट्र शासन" head line
SIZE_REFS = 11.5
SIZE_HF = 10.0                     # running header + footer
SIZE_COPY = 12.0                   # प्रति (tight pitch makes it read smaller)

SUBJECT_X = 324.0                  # subject box: x 324..543
SUBJECT_W = PAGE_W - MARGIN_R - SUBJECT_X      # 219
SUBJECT_RULE_W = 175.0             # thin rule under the last subject line
LETTERHEAD_GAP = 19.0

READ_NUM_X = 108.0                 # वाचा item number x
READ_TEXT_X = 126.0                # वाचा item text + continuation x
COPY_NUM_X = 117.0                 # प्रति number x
COPY_TEXT_X = 135.0                # प्रति text + continuation x

HEADER_Y = 21.5
RULE_X0, RULE_X1 = 70.6, 541.5
RULE_THIN_Y, RULE_THIN_H = 37.6, 0.72
RULE_THICK_Y, RULE_THICK_H = 39.0, 2.16
FOOTER_Y = 758.0
RIGHT_X = PAGE_W - MARGIN_R        # 543

SIGN_CX = 405.0                    # signature name/designation centre axis

DEV_DIGITS = str.maketrans("0123456789", "०१२३४५६७८९")

L10N = {
    "mr": {
        "govt": "महाराष्ट्र शासन",
        "grno": "शासन निर्णय क्रमांक:- {no},",
        "addr1": "मादाम कामा रोड, हुतात्मा राजगुरु चौक,",
        "addr2": "मंत्रालय, मुंबई-400 032,",
        "date": "दिनांक:- {date}",
        "read": "वाचा :-",
        "preamble": "प्रस्तावना:-",
        "resolution": "शासन निर्णय :-",
        "avail": "सदर शासन निर्णय महाराष्ट्र शासनाच्या "
                 "[www.maharashtra.gov.in](https://www.maharashtra.gov.in) या "
                 "संकेतस्थळावर उपलब्ध करण्यात आला असून त्याचा संगणक सांकेतांक "
                 "{code} असा आहे. हा आदेश डिजीटल स्वाक्षरीने साक्षांकित करुन "
                 "काढण्यात येत आहे.",
        "governor": "महाराष्ट्राचे राज्यपाल यांच्या आदेशानुसार व नावाने.",
        "sign_suffix": ", महाराष्ट्र शासन",
        "copy": "प्रत,",
        "hdr_prefix": "शासन निर्णय क्रमांकः",
        "footer": "पृष्ठ {total} पैकी {n}",
        "fin_clause": "या निर्णयापोटी होणारा खर्च मंजूर अर्थसंकल्पीय तरतूद, "
                      "लेखाशीर्ष {head} मधून सन {fy} या आर्थिक वर्षात "
                      "भागविण्यात यावा.",
        "fin_marker": "लेखाशीर्ष",
    },
    "en": {
        "govt": "GOVERNMENT OF MAHARASHTRA",
        "grno": "Government Resolution No.:- {no},",
        "addr1": "Madam Cama Road, Hutatma Rajguru Chowk,",
        "addr2": "Mantralaya, Mumbai - 400 032,",
        "date": "Date:- {date}",
        "read": "Read :-",
        "preamble": "Preamble :-",
        "resolution": "Government Resolution :-",
        "avail": "This Government Resolution is available on the Government of "
                 "Maharashtra website "
                 "[www.maharashtra.gov.in](https://www.maharashtra.gov.in), "
                 "with Unique Code {code}. This order is issued attested by "
                 "digital signature.",
        "governor": "By order and in the name of the Governor of Maharashtra.",
        "sign_suffix": ", Government of Maharashtra",
        "copy": "Copy to,",
        "hdr_prefix": "Government Resolution No.:",
        "footer": "Page {n} of {total}",
        "fin_clause": "The expenditure on account of this decision shall be met "
                      "from the sanctioned budgetary provision under Budget "
                      "Head {head} for the financial year {fy}.",
        "fin_marker": "Budget Head",
    },
}


def to_devanagari_digits(s: str) -> str:
    return (s or "").translate(DEV_DIGITS)


def _marathi_date(d: date_cls) -> str:
    return f"{d.day} {MARATHI_MONTHS[d.month - 1]}, {d.year}"


def _english_date(d: date_cls) -> str:
    return f"{d.day} {d:%B}, {d.year}"


def _normalize(form: dict) -> dict:
    """Fill every key the renderer reads, so both the structured GRForm and a
    best-effort parse of a hand-typed draft render without exceptions."""
    f = dict(form or {})
    blob = " ".join(str(f.get(k) or "") for k in
                    ("subject", "background", "department"))
    f.setdefault("lang", "mr" if re.search(r"[ऀ-ॿ]", blob) else "en")
    lang = f["lang"] = f["lang"] if f["lang"] in L10N else "en"

    d = f.get("date")
    if isinstance(d, str):
        try:
            d = date_cls.fromisoformat(d[:10])
        except ValueError:
            d = None
    if not isinstance(d, date_cls):
        d = date_cls.today()
    f["date"] = d

    for k in ("subject", "department", "gr_number", "background",
              "budget_head", "signatory", "designation", "topic"):
        f[k] = str(f.get(k) or "").strip()
    f.setdefault("fy", financial_year(d))
    f.setdefault("code", unique_code(d, f["subject"]))

    refs = []
    for r in f.get("references") or []:
        text = (r.get("text") if isinstance(r, dict) else str(r)).strip()
        if text:
            refs.append(text)
    f["references"] = refs

    f["clauses"] = [str(c).strip() for c in f.get("clauses") or [] if str(c).strip()]
    # Keep the rendered GR consistent with the canonical build_draft() text:
    # the financial-sanction clause is auto-appended there, so append the same
    # wording here unless the clause list already carries a budget head.
    t = L10N[lang]
    if f["budget_head"] and not any(t["fin_marker"] in c for c in f["clauses"]):
        f["clauses"].append(t["fin_clause"].format(head=f["budget_head"], fy=f["fy"]))

    if not f.get("preamble_paras"):
        f["preamble_paras"] = [p.strip() for p in
                               re.split(r"\n\s*\n", f["background"]) if p.strip()]
    f["copy_to"] = [str(c).strip() for c in f.get("copy_to") or [] if str(c).strip()]
    return f


class _GRPDF(FPDF):
    """US-Letter page with the gold specimen's running header (pages 2+, file
    number bold, thin-over-thick rule) and right-aligned Devanagari footer.
    Page totals need two passes: pass 1 with total_pages=None (footer off) just
    counts pages; pass 2 renders the real footer."""

    MARKDOWN_LINK_COLOR = "#0000ee"

    def __init__(self, lang: str, gr_number: str, total_pages=None):
        super().__init__(unit="pt", format="letter")
        self.lang = lang
        self.T = to_devanagari_digits if lang == "mr" else (lambda s: s)
        self.hdr_prefix = L10N[lang]["hdr_prefix"]
        self.hdr_no = self.T(gr_number)
        self.total_pages = total_pages
        for fam, style, fname in (("noto", "", "NotoSans-Regular.ttf"),
                                  ("noto", "B", "NotoSans-Bold.ttf"),
                                  ("notodev", "", "NotoSansDevanagari-Regular.ttf"),
                                  ("notodev", "B", "NotoSansDevanagari-Bold.ttf")):
            self.add_font(fam, style, str(FONT_DIR / fname))
        self.set_fallback_fonts(["notodev"])
        self.set_text_shaping(True)
        self.set_margins(MARGIN_L, CONTENT_TOP, MARGIN_R)
        self.set_auto_page_break(True, BREAK_MARGIN)

    def header(self):
        if self.page_no() == 1:
            return
        _set_font(self, "", SIZE_HF)
        self.set_xy(MARGIN_L, HEADER_Y)
        self.cell(self.get_string_width(self.hdr_prefix + " "), 12,
                  self.hdr_prefix + " ")
        _set_font(self, "B", SIZE_HF)
        self.cell(0, 12, self.hdr_no)
        self.set_fill_color(0, 0, 0)
        self.rect(RULE_X0, RULE_THIN_Y, RULE_X1 - RULE_X0, RULE_THIN_H, style="F")
        self.rect(RULE_X0, RULE_THICK_Y, RULE_X1 - RULE_X0, RULE_THICK_H, style="F")
        _set_font(self)
        self.set_y(CONTENT_TOP)

    def footer(self):
        if self.page_no() == 1 or self.total_pages is None:
            return
        _set_font(self, "", SIZE_HF)
        txt = self.T(L10N[self.lang]["footer"].format(total=self.total_pages,
                                                      n=self.page_no()))
        w = self.get_string_width(txt) + 2
        self.set_xy(RIGHT_X - w, FOOTER_Y)
        self.cell(w, 12, txt, align="R")


# ---------- layout primitives ----------

def _set_font(pdf, style="", size=SIZE_BODY):
    """set_font with the fallback-poisoning defeated.

    fpdf2 2.8.7: rendering a fragment whose glyphs came from a fallback font
    leaves pdf.current_font pointing at the fallback while font_family still
    reads "noto", and set_font("noto") short-circuits without repairing it.
    Everything parsed afterwards then resolves glyph presence against the
    WRONG cmap: Latin text vanishes from Devanagari paragraphs, and markdown
    mode (which reassigns frag.font from font_family) blanks Devanagari.
    Forcing current_font after set_font makes every primitive parse against
    the true base font."""
    pdf.set_font("noto", style, size)
    pdf.current_font = pdf.fonts["noto" + style]


def _need(pdf, h):
    """Force the page break BEFORE a two-part row (number cell + hanging text)
    so the number can never strand on the previous page."""
    if pdf.y + h > PAGE_H - BREAK_MARGIN:
        pdf.add_page()


def _subject_block(pdf, subject: str, topic: str):
    _set_font(pdf, "B")
    pdf.set_left_margin(SUBJECT_X)
    pdf.set_xy(SUBJECT_X, CONTENT_TOP)
    pdf.multi_cell(SUBJECT_W, LINE_H_TIGHT, subject, align="J")
    pdf.set_fill_color(0, 0, 0)
    pdf.rect(SUBJECT_X, pdf.get_y() + 1.5, SUBJECT_RULE_W, 0.4, style="F")
    if topic and topic != subject:
        pdf.set_y(pdf.get_y() + 3.5)
        pdf.multi_cell(SUBJECT_W, LINE_H_TIGHT, topic, align="J")
    pdf.set_left_margin(MARGIN_L)
    pdf.set_x(MARGIN_L)


def _letterhead(pdf, lines):
    pdf.set_y(pdf.get_y() + LETTERHEAD_GAP)
    for i, line in enumerate(lines):
        _set_font(pdf, "B", SIZE_TITLE if i == 0 else SIZE_BODY)
        pdf.set_x(MARGIN_L)
        pdf.multi_cell(BODY_W, LINE_H_TIGHT, line, align="C")
    _set_font(pdf)


def _label_line(pdf, label: str, gap_before=SECTION_GAP):
    pdf.set_y(pdf.get_y() + gap_before)
    _need(pdf, LINE_H_BODY)
    _set_font(pdf, "B")
    pdf.set_x(MARGIN_L)
    pdf.cell(0, LINE_H_BODY, label)
    pdf.ln(LINE_H_BODY)
    _set_font(pdf)


def _hanging_item(pdf, num: str, text: str, num_x: float, text_x: float,
                  h: float, size: float):
    _need(pdf, h)
    _set_font(pdf, "", size)
    y0 = pdf.get_y()
    pdf.set_xy(num_x, y0)
    pdf.cell(text_x - num_x, h, num)
    pdf.set_left_margin(text_x)
    pdf.set_xy(text_x, y0)
    pdf.multi_cell(RIGHT_X - text_x, h, text, align="L")
    pdf.set_left_margin(MARGIN_L)
    pdf.set_x(MARGIN_L)


def _body_para(pdf, text: str, first_indent=0.0, align="J", markdown=False,
               gap_after=PARA_GAP):
    _set_font(pdf)
    pdf.set_x(MARGIN_L)
    if first_indent:
        # fpdf2 has no native first-line indent for multi_cell; an EM-space run
        # survives shaping and is not stretched away by justification.
        n = max(1, round(first_indent / pdf.get_string_width(" ")))
        text = " " * n + text
    pdf.multi_cell(BODY_W, LINE_H_BODY, text, align=align, markdown=markdown)
    if gap_after:
        pdf.set_y(pdf.get_y() + gap_after)


def _signature(pdf, lang: str, signatory: str, designation: str):
    t = L10N[lang]
    _need(pdf, 5 * LINE_H_BODY + 60)
    pdf.set_y(pdf.get_y() + SECTION_GAP)
    _set_font(pdf)
    pdf.set_x(MARGIN_L)
    pdf.multi_cell(BODY_W, LINE_H_BODY, t["governor"], align="C")
    pdf.set_y(pdf.get_y() + 60)          # space where the stamp goes at issuance
    for line, style in ((f"( {signatory} )", "B"),
                        (designation + t["sign_suffix"], "")):
        _set_font(pdf, style)
        w = pdf.get_string_width(line)
        pdf.set_xy(SIGN_CX - w / 2, pdf.get_y())
        pdf.cell(w, LINE_H_BODY, line)
        pdf.ln(LINE_H_BODY)


def _copy_to(pdf, lang: str, entries):
    if not entries:
        return
    _label_line(pdf, L10N[lang]["copy"], gap_before=SECTION_GAP)
    T = to_devanagari_digits if lang == "mr" else (lambda s: s)
    for i, entry in enumerate(entries, 1):
        _hanging_item(pdf, T(f"{i})"), T(entry), COPY_NUM_X, COPY_TEXT_X,
                      LINE_H_TIGHT, SIZE_COPY)


# ---------- document flow ----------

def _render_gr(form: dict, total_pages=None) -> _GRPDF:
    f = _normalize(form)
    lang, t = f["lang"], L10N[f["lang"]]
    pdf = _GRPDF(lang, f["gr_number"], total_pages)
    T = pdf.T
    pdf.add_page()

    _subject_block(pdf, T(f["subject"]), T(f["topic"]))
    date_str = _marathi_date(f["date"]) if lang == "mr" else _english_date(f["date"])
    _letterhead(pdf, [T(line) for line in
                      (t["govt"], f["department"],
                       t["grno"].format(no=f["gr_number"]),
                       t["addr1"], t["addr2"],
                       t["date"].format(date=date_str))])

    # वाचा label shares the first reference's baseline, as in the gold specimen
    pdf.set_y(pdf.get_y() + SECTION_GAP)
    _need(pdf, LINE_H_TIGHT)
    _set_font(pdf, "B")
    label_w = pdf.get_string_width(t["read"])
    pdf.set_xy(MARGIN_L, pdf.get_y())
    pdf.cell(label_w, LINE_H_TIGHT, t["read"])
    num_x = max(READ_NUM_X, MARGIN_L + label_w + 6)
    text_x = num_x + (READ_TEXT_X - READ_NUM_X)
    refs = f["references"] or ["—"]
    for i, ref in enumerate(refs, 1):
        n = T(f"{i})") if f["references"] else ""
        if i == 1:
            _hanging_item(pdf, n, T(ref), num_x, text_x,
                          LINE_H_TIGHT, SIZE_REFS)
        else:
            _hanging_item(pdf, n, T(ref), READ_NUM_X, READ_TEXT_X,
                          LINE_H_TIGHT, SIZE_REFS)

    _label_line(pdf, t["preamble"])
    for para in f["preamble_paras"]:
        _body_para(pdf, T(para), first_indent=FIRST_INDENT)

    _label_line(pdf, t["resolution"])
    for i, clause in enumerate(f["clauses"], 1):
        _body_para(pdf, T(f"{i}. {clause}"))
    # availability paragraph = the final numbered clause of a real GR, with the
    # blue clickable URL; the only paragraph rendered in markdown mode
    avail = t["avail"].format(code=T(f["code"]))
    _body_para(pdf, T(f"{len(f['clauses']) + 1}. ") + avail, markdown=True)

    _signature(pdf, lang, T(f["signatory"]), T(f["designation"]))
    _copy_to(pdf, lang, f["copy_to"])
    return pdf


def gr_pdf(form: dict) -> bytes:
    """Authentic-format GR PDF (two passes so the footer can carry the real
    page total in Devanagari digits — the {nb} alias cannot)."""
    n = _render_gr(form, total_pages=None).page_no()
    pdf = _render_gr(form, total_pages=n)
    assert pdf.page_no() == n, "page count changed between footer passes"
    return bytes(pdf.output())


# ---------- DOCX parity ----------

def _bold_cs(run):
    """Bold that actually applies to Devanagari: w:b alone does not embolden
    complex-script runs — w:bCs must be set too."""
    from docx.oxml.ns import qn
    run.font.bold = True
    rpr = run._element.get_or_add_rPr()
    if rpr.find(qn("w:bCs")) is None:
        el = rpr.makeelement(qn("w:bCs"), {})
        rpr.append(el)


def _run_sz_cs(run, size_pt: float):
    """Complex-script size (w:szCs) — Word sizes Devanagari from szCs, which
    python-docx's font.size does not set."""
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = rpr.makeelement(qn("w:szCs"), {})
        rpr.append(szcs)
    szcs.set(qn("w:val"), str(int(size_pt * 2)))


def _para(container, text="", *, align=None, bold=False, size=SIZE_BODY,
          left_indent=None, first_line_indent=None, line_exact=LINE_H_BODY,
          space_after=0.0, tabs=()):
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt
    p = container.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if left_indent is not None:
        pf.left_indent = Pt(left_indent)
    if first_line_indent is not None:
        pf.first_line_indent = Pt(first_line_indent)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line_exact)
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    for x in tabs:
        pf.tab_stops.add_tab_stop(Pt(x))
    if text:
        run = p.add_run(text)
        _run_sz_cs(run, size)
        if bold:
            _bold_cs(run)
    return p


def _p_border_bottom(p, val="single", sz="6"):
    from docx.oxml.ns import qn
    ppr = p._element.get_or_add_pPr()
    pbdr = ppr.makeelement(qn("w:pBdr"), {})
    bottom = ppr.makeelement(qn("w:bottom"), {})
    bottom.set(qn("w:val"), val)
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom)
    ppr.append(pbdr)


def _add_field(p, instr: str, cached: str = "1"):
    """PAGE / NUMPAGES field with a cached value so LibreOffice renders it
    without a manual refresh."""
    from docx.oxml.ns import qn
    r1 = p.add_run()._element
    fld = r1.makeelement(qn("w:fldChar"), {})
    fld.set(qn("w:fldCharType"), "begin")
    r1.append(fld)
    r2 = p.add_run()._element
    it = r2.makeelement(qn("w:instrText"), {})
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2.append(it)
    r3 = p.add_run()._element
    sep = r3.makeelement(qn("w:fldChar"), {})
    sep.set(qn("w:fldCharType"), "separate")
    r3.append(sep)
    p.add_run(cached)
    r4 = p.add_run()._element
    end = r4.makeelement(qn("w:fldChar"), {})
    end.set(qn("w:fldCharType"), "end")
    r4.append(end)


def _add_hyperlink(p, text: str, url: str, size=SIZE_BODY):
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = p._element.makeelement(qn("w:hyperlink"), {})
    hl.set(qn("r:id"), r_id)
    run = p.add_run(text)
    _run_sz_cs(run, size)
    run.font.underline = True
    from docx.shared import RGBColor
    run.font.color.rgb = RGBColor(0x00, 0x00, 0xEE)
    p._element.remove(run._element)
    hl.append(run._element)
    p._element.append(hl)


def gr_docx(form: dict) -> bytes:
    """Authentic-format GR as DOCX (Nirmala UI — Word shapes the Devanagari).
    Same geometry as the PDF; accepted deviation: PAGE/NUMPAGES field digits
    render in ASCII (fields cannot emit Devanagari digits), all literal text
    still carries Devanagari digits. The PDF is the pixel-faithful artifact;
    this is the editable companion."""
    import io
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
    from docx.oxml.ns import qn
    from docx.shared import Pt

    f = _normalize(form)
    lang, t = f["lang"], L10N[f["lang"]]
    T = to_devanagari_digits if lang == "mr" else (lambda s: s)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Nirmala UI"
    style.font.size = Pt(SIZE_BODY)
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Nirmala UI")
    rfonts.set(qn("w:cs"), "Nirmala UI")

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(PAGE_W), Pt(PAGE_H)
    sec.left_margin, sec.right_margin = Pt(MARGIN_L), Pt(MARGIN_R)
    sec.top_margin, sec.bottom_margin = Pt(CONTENT_TOP), Pt(BREAK_MARGIN)
    sec.header_distance, sec.footer_distance = Pt(HEADER_Y), Pt(PAGE_H - FOOTER_Y - 14)
    sec.different_first_page_header_footer = True

    # running header (pages 2+): file-number line, bold number, double rule
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run(t["hdr_prefix"] + " ")
    _run_sz_cs(r, SIZE_HF)
    rb = hp.add_run(T(f["gr_number"]))
    _run_sz_cs(rb, SIZE_HF)
    _bold_cs(rb)
    _p_border_bottom(hp, val="thinThickSmallGap", sz="12")
    # footer (pages 2+): पृष्ठ <total> पैकी <n>, right-aligned
    fp = sec.footer.paragraphs[0]
    fp.alignment = AL.RIGHT
    if lang == "mr":
        r = fp.add_run("पृष्ठ ")
        _run_sz_cs(r, SIZE_HF)
        _add_field(fp, "NUMPAGES")
        r = fp.add_run(" पैकी ")
        _run_sz_cs(r, SIZE_HF)
        _add_field(fp, "PAGE")
    else:
        r = fp.add_run("Page ")
        _run_sz_cs(r, SIZE_HF)
        _add_field(fp, "PAGE")
        r = fp.add_run(" of ")
        _run_sz_cs(r, SIZE_HF)
        _add_field(fp, "NUMPAGES")
    # first page: empty header/footer
    sec.first_page_header.paragraphs[0].text = ""
    sec.first_page_footer.paragraphs[0].text = ""

    # subject block (top-right), rule under the subject proper
    sp = _para(doc, T(f["subject"]), align=AL.JUSTIFY, bold=True,
               left_indent=SUBJECT_X - MARGIN_L, line_exact=LINE_H_TIGHT)
    _p_border_bottom(sp)
    if f["topic"] and f["topic"] != f["subject"]:
        _para(doc, T(f["topic"]), align=AL.JUSTIFY, bold=True,
              left_indent=SUBJECT_X - MARGIN_L, line_exact=LINE_H_TIGHT)

    # letterhead
    date_str = _marathi_date(f["date"]) if lang == "mr" else _english_date(f["date"])
    head = [(t["govt"], SIZE_TITLE), (T(f["department"]), SIZE_BODY),
            (T(t["grno"].format(no=f["gr_number"])), SIZE_BODY),
            (t["addr1"], SIZE_BODY), (T(t["addr2"]), SIZE_BODY),
            (T(t["date"].format(date=date_str)), SIZE_BODY)]
    for i, (line, size) in enumerate(head):
        _para(doc, line, align=AL.CENTER, bold=True, size=size,
              line_exact=LINE_H_TIGHT,
              space_after=PARA_GAP if i == len(head) - 1 else 0)

    # वाचा — label + first ref on one line, hanging indents
    refs = [T(r["text"] if isinstance(r, dict) else str(r))
            for r in f["references"]] or ["—"]
    first = _para(doc, align=AL.LEFT, line_exact=LINE_H_TIGHT,
                  left_indent=READ_TEXT_X - MARGIN_L,
                  first_line_indent=-(READ_TEXT_X - MARGIN_L),
                  tabs=(READ_NUM_X - MARGIN_L, READ_TEXT_X - MARGIN_L))
    lr = first.add_run(t["read"])
    _run_sz_cs(lr, SIZE_BODY)
    _bold_cs(lr)
    rr = first.add_run("\t" + (T("1)") if f["references"] else "") + "\t" + refs[0])
    _run_sz_cs(rr, SIZE_REFS)
    for i, ref in enumerate(refs[1:], 2):
        p = _para(doc, line_exact=LINE_H_TIGHT,
                  left_indent=READ_TEXT_X - MARGIN_L,
                  first_line_indent=-(READ_TEXT_X - READ_NUM_X),
                  tabs=(READ_TEXT_X - MARGIN_L,))
        run = p.add_run(T(f"{i})") + "\t" + ref)
        _run_sz_cs(run, SIZE_REFS)

    # प्रस्तावना
    _para(doc, t["preamble"], bold=True, space_after=0,
          line_exact=LINE_H_BODY)
    for para in f["preamble_paras"]:
        _para(doc, T(para), align=AL.JUSTIFY, first_line_indent=FIRST_INDENT,
              space_after=PARA_GAP)

    # शासन निर्णय
    _para(doc, t["resolution"], bold=True, line_exact=LINE_H_BODY)
    for i, clause in enumerate(f["clauses"], 1):
        _para(doc, T(f"{i}. {clause}"), align=AL.JUSTIFY, space_after=PARA_GAP)
    avail_pre, avail_post = t["avail"].split(
        "[www.maharashtra.gov.in](https://www.maharashtra.gov.in)")
    ap = _para(doc, align=AL.JUSTIFY, space_after=PARA_GAP)
    run = ap.add_run(T(f"{len(f['clauses']) + 1}. ") + T(avail_pre))
    _run_sz_cs(run, SIZE_BODY)
    _add_hyperlink(ap, "www.maharashtra.gov.in", "https://www.maharashtra.gov.in")
    run = ap.add_run(T(avail_post.format(code=T(f["code"]))))
    _run_sz_cs(run, SIZE_BODY)

    # signature — no fabricated stamp; centred on the SIGN_CX axis
    _para(doc, t["governor"], align=AL.CENTER, space_after=60)
    off = 2 * (SIGN_CX - (MARGIN_L + (PAGE_W - MARGIN_R)) / 2)
    _para(doc, T(f"( {f['signatory']} )"), align=AL.CENTER, bold=True,
          left_indent=off)
    _para(doc, T(f["designation"] + t["sign_suffix"]), align=AL.CENTER,
          left_indent=off, space_after=PARA_GAP)

    # प्रति
    if f["copy_to"]:
        _para(doc, t["copy"], bold=True, line_exact=LINE_H_BODY)
        for i, entry in enumerate(f["copy_to"], 1):
            p = _para(doc, line_exact=LINE_H_TIGHT,
                      left_indent=COPY_TEXT_X - MARGIN_L,
                      first_line_indent=-(COPY_TEXT_X - COPY_NUM_X),
                      tabs=(COPY_TEXT_X - MARGIN_L,))
            run = p.add_run(T(f"{i})") + "\t" + T(entry))
            _run_sz_cs(run, SIZE_COPY)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    # Calibration gate: render the SAMPLE fixture in both languages, rasterize
    # alongside the gold-standard pages for side-by-side eyeballing.
    sys.stdout.reconfigure(encoding="utf-8")
    from generate import SAMPLE
    import fitz

    out = ROOT / "data" / "gr_layout_check"
    out.mkdir(exist_ok=True)
    copy_to = ["मा. मुख्यमंत्री यांचे प्रधान सचिव, मंत्रालय, मुंबई",
               "मा. उपमुख्यमंत्री (वित्त व नियोजन) यांचे सचिव, मंत्रालय, मुंबई",
               "मा. मंत्री (उच्च व तंत्र शिक्षण) यांचे खाजगी सचिव",
               "मा. मुख्य सचिव, महाराष्ट्र शासन, मंत्रालय, मुंबई.",
               "अपर मुख्य सचिव, वित्त विभाग, मंत्रालय, मुंबई",
               "प्रधान सचिव, उच्च व तंत्र शिक्षण विभाग, मंत्रालय, मुंबई",
               "संचालक, तंत्रशिक्षण संचालनालय, महाराष्ट्र राज्य, मुंबई",
               "जिल्हाधिकारी, सर्व",
               "महालेखापाल (लेखा व अनुज्ञेयता)-1/2 महाराष्ट्र, मुंबई/नागपूर",
               "महालेखापाल (लेखापरीक्षा)-1/2, महाराष्ट्र, मुंबई/नागपूर",
               "सर्व शासकीय तंत्रनिकेतनांचे प्राचार्य",
               "उच्च व तंत्र शिक्षण विभागातील सर्व कार्यासने",
               "निवड नस्ती, ताशि-2."]
    for lang in ("mr", "en"):
        form = dict(SAMPLE[lang], date=date_cls(2026, 8, 4),
                    topic="", copy_to=copy_to if lang == "mr" else
                    ["Principal Secretary, H&TE Dept, Mantralaya, Mumbai",
                     "Director of Technical Education, Mumbai",
                     "Accountant General (A&E)-1/2, Maharashtra",
                     "Select file, TE-2."])
        data = gr_pdf(form)
        (out / f"selftest_{lang}.pdf").write_bytes(data)
        (out / f"selftest_{lang}.docx").write_bytes(gr_docx(form))
        doc = fitz.open(stream=data, filetype="pdf")
        for i, page in enumerate(doc, 1):
            page.get_pixmap(dpi=140).save(out / f"gen_{lang}_p{i}.png")
        print(f"[{lang}] {doc.page_count} pages -> {out}\\gen_{lang}_p*.png + .docx")
    gold = fitz.open(ROOT / "data" / "raw_pdfs" / "202402271742452721.pdf")
    for i, page in enumerate(gold, 1):
        page.get_pixmap(dpi=140).save(out / f"gold_p{i}.png")
    print(f"gold pages -> {out}\\gold_p*.png — compare side by side")
